"""
Récupération automatique de la mesure de fin CIT depuis THERMAL_CYCLING.

Source prioritaire : *.xlsm feuille data_traitee (dernière valeur du canal SN).
Fallback : *acceptance*.XLS OptoTest (mapping SN ↔ CHn via le nom de fichier).

Valeur retournée = abs(dernière mesure), arrondie à 3 décimales.
"""

from __future__ import annotations

import os
import re
from typing import Optional

import Log

THERMAL_FOLDER_NAME = "THERMAL_CYCLING"
DATA_SHEET_NAME = "data_traitee"
HEADER_ROW_INDEX = 2  # 0-based → ligne 3 Excel
CIT_TAG_ALIASES = ("CIT", "FIN_CIT", "fin_cit")


def _normalize_sn(sn_value) -> str:
    """Normalise un SN : 144, 0144, 25-20-01, SN25-20-01, 25/47/02."""
    if sn_value is None:
        return ""
    sn = str(sn_value).strip()
    if not sn:
        return ""
    sn = re.sub(r"(?i)^s/?n\s*", "", sn).strip()
    sn = sn.replace("/", "-")
    if "-" in sn:
        parts = [p.strip() for p in sn.split("-") if p.strip()]
        return "-".join(p.lstrip("0") or "0" for p in parts)
    return sn.lstrip("0") or "0"


def _extract_sn_from_header(text: str) -> str:
    """Extrait le SN depuis un en-tête type 'RETA-AGS24.134B SN25-20-01'."""
    if text is None:
        return ""
    text = str(text).strip()
    if not text:
        return ""
    # SN avec tirets/slashs ou numérique simple en fin de chaîne
    match = re.search(
        r"(?i)(?:s/?n\s*)?(\d+(?:[-/]\d+)*)\s*$",
        text,
    )
    if match:
        return _normalize_sn(match.group(1))
    return ""


def _normalize_ref(ref_value) -> str:
    if ref_value is None:
        return ""
    ref = str(ref_value).strip().upper()
    # Harmoniser variantes d'encodage / caractères spéciaux courants
    ref = ref.replace("°", "").replace("º", "")
    ref = re.sub(r"\s+", "", ref)
    return ref


def _format_cit_value(raw_value) -> Optional[str]:
    """abs + arrondi 3 décimales, format point (cohérent mesures Excel)."""
    try:
        if raw_value is None or raw_value == "":
            return None
        if isinstance(raw_value, str):
            text = raw_value.strip().replace(",", ".")
            if not text:
                return None
            number = float(text)
        else:
            number = float(raw_value)
        if number != number:  # NaN
            return None
        return f"{abs(number):.3f}"
    except (TypeError, ValueError):
        return None


def _header_matches(header, sn: str, ref: str) -> bool:
    """True si l'en-tête (ex. 'RCT°-AGS23.157 SN144') correspond à ref+SN."""
    if header is None:
        return False
    text = str(header).strip()
    if not text:
        return False

    sn_norm = _normalize_sn(sn)
    if not sn_norm:
        return False

    header_sn = _extract_sn_from_header(text)
    if not header_sn or header_sn != sn_norm:
        return False

    if not ref:
        # Sans référence fournie, le SN seul suffit (fichier mono-réf typique)
        return True

    sn_match = re.search(r"(?i)(?:s/?n\s*)?\d+(?:[-/]\d+)*\s*$", text)
    header_ref_part = text[: sn_match.start()].strip() if sn_match else text
    header_ref = _normalize_ref(header_ref_part)
    target_ref = _normalize_ref(ref)
    if not header_ref:
        return True

    # Match souple : l'un contient l'autre (ex. 23.157 vs RCT-AGS23.157C)
    return target_ref in header_ref or header_ref in target_ref


def _is_temp_excel(filename: str) -> bool:
    name = os.path.basename(filename)
    return name.startswith("~$")


def _list_thermal_files(thermal_dir: str, extensions: tuple) -> list:
    if not thermal_dir or not os.path.isdir(thermal_dir):
        return []
    files = []
    for name in os.listdir(thermal_dir):
        if _is_temp_excel(name):
            continue
        lower = name.lower()
        if any(lower.endswith(ext) for ext in extensions):
            files.append(os.path.join(thermal_dir, name))
    files.sort()
    return files


def _find_last_numeric_in_column(rows, col_index: int):
    """Parcourt les lignes de bas en haut pour la dernière valeur numérique."""
    for row in reversed(rows):
        if row is None or col_index >= len(row):
            continue
        cell = row[col_index]
        formatted = _format_cit_value(cell)
        if formatted is not None:
            return cell
    return None


def _extract_ref_from_header(text: str) -> str:
    """Extrait la référence depuis un en-tête 'RCT°-AGS23.157 SN144'."""
    if text is None:
        return ""
    text = str(text).strip()
    if not text:
        return ""
    sn_match = re.search(r"(?i)(?:s/?n\s*)?\d+(?:[-/]\d+)*\s*$", text)
    if not sn_match:
        return _normalize_ref(text)
    return _normalize_ref(text[: sn_match.start()].strip())


def _cit_lookup_key(sn: str, ref: str) -> tuple:
    return (_normalize_sn(sn), _normalize_ref(ref))


def _lookup_cit_in_index(index: dict, sn: str, ref: str) -> Optional[str]:
    if not index:
        return None
    sn_norm = _normalize_sn(sn)
    ref_norm = _normalize_ref(ref)
    if not sn_norm:
        return None

    exact = index.get((sn_norm, ref_norm))
    if exact is not None:
        return exact

    # Match souple référence + SN
    for (idx_sn, idx_ref), value in index.items():
        if idx_sn != sn_norm:
            continue
        if not ref_norm or not idx_ref:
            return value
        if ref_norm in idx_ref or idx_ref in ref_norm:
            return value
    return None


def _parse_xlsm_rows_to_index(rows) -> dict:
    """Construit un index {(sn, ref): valeur} depuis les lignes data_traitee."""
    index = {}
    if len(rows) <= HEADER_ROW_INDEX:
        return index

    headers = rows[HEADER_ROW_INDEX]
    data_rows = rows[HEADER_ROW_INDEX + 1 :]

    for col_index, header in enumerate(headers):
        if header is None:
            continue
        sn_h = _extract_sn_from_header(str(header))
        if not sn_h:
            continue
        ref_h = _extract_ref_from_header(str(header))
        raw = _find_last_numeric_in_column(data_rows, col_index)
        formatted = _format_cit_value(raw)
        if formatted is None:
            continue
        key = (sn_h, ref_h)
        if key not in index:
            index[key] = formatted
    return index


def _parse_xlsm_file_to_index(path: str) -> dict:
    """Lit un .xlsm/.xlsx et retourne toutes les mesures de fin du fichier."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            if DATA_SHEET_NAME not in wb.sheetnames:
                return {}
            ws = wb[DATA_SHEET_NAME]
            rows = list(ws.iter_rows(values_only=True))
            return _parse_xlsm_rows_to_index(rows)
        finally:
            wb.close()
    except ImportError:
        return _parse_xlsm_file_to_index_com(path)
    except Exception as e:
        Log.Verbose(f"[CIT] openpyxl index {os.path.basename(path)}: {e}")
        return _parse_xlsm_file_to_index_com(path)


def _parse_xlsm_file_to_index_com(path: str) -> dict:
    """Indexe data_traitee via Excel COM (une ouverture par fichier)."""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return {}

    app = None
    wb = None
    index = {}
    try:
        try:
            pythoncom.CoInitialize()
        except Exception:
            pass
        try:
            app = win32com.client.DispatchEx("Excel.Application")
        except Exception:
            app = win32com.client.Dispatch("Excel.Application")
        try:
            app.Visible = 0
            app.DisplayAlerts = 0
        except Exception:
            pass

        wb = app.Workbooks.Open(path, ReadOnly=True)
        try:
            ws = wb.Sheets(DATA_SHEET_NAME)
        except Exception:
            return {}

        used_cols = ws.UsedRange.Columns.Count
        used_rows = ws.UsedRange.Rows.Count
        header_row = HEADER_ROW_INDEX + 1

        for col in range(1, used_cols + 1):
            header = ws.Cells(header_row, col).Value
            if header is None:
                continue
            header_str = str(header)
            sn_h = _extract_sn_from_header(header_str)
            if not sn_h:
                continue
            ref_h = _extract_ref_from_header(header_str)
            raw = None
            for row in range(used_rows, header_row, -1):
                cell = ws.Cells(row, col).Value
                if _format_cit_value(cell) is not None:
                    raw = cell
                    break
            formatted = _format_cit_value(raw)
            if formatted is None:
                continue
            key = (sn_h, ref_h)
            if key not in index:
                index[key] = formatted
        return index
    except Exception as e:
        Log.Verbose(f"[CIT] COM index {os.path.basename(path)}: {e}")
        return {}
    finally:
        try:
            if wb is not None:
                wb.Close(False)
        except Exception:
            pass
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass


def _parse_acceptance_file_to_index(path: str) -> dict:
    """Indexe toutes les mesures de fin d'un fichier acceptance OptoTest."""
    channels = _parse_acceptance_filename_channels(path)
    if not channels:
        return {}

    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return {}

    app = None
    wb = None
    index = {}
    try:
        try:
            pythoncom.CoInitialize()
        except Exception:
            pass
        try:
            app = win32com.client.DispatchEx("Excel.Application")
        except Exception:
            app = win32com.client.Dispatch("Excel.Application")
        try:
            app.Visible = 0
            app.DisplayAlerts = 0
        except Exception:
            pass

        wb = app.Workbooks.Open(path, ReadOnly=True)
        ws = wb.Sheets(1)
        used_rows = ws.UsedRange.Rows.Count

        for idx, ch in enumerate(channels):
            if ch is None:
                continue
            target_col = 6 + idx
            raw = None
            for row in range(used_rows, 6, -1):
                cell = ws.Cells(row, target_col).Value
                if _format_cit_value(cell) is not None:
                    raw = cell
                    break
            formatted = _format_cit_value(raw)
            if formatted is None:
                continue
            sn_h = ch["sn"]
            ref_h = _normalize_ref(ch.get("ref_hint") or "")
            key = (sn_h, ref_h)
            if key not in index:
                index[key] = formatted
        return index
    except Exception as e:
        Log.Verbose(f"[CIT] acceptance index {os.path.basename(path)}: {e}")
        return {}
    finally:
        try:
            if wb is not None:
                wb.Close(False)
        except Exception:
            pass
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass


def _resolve_thermal_dir(work_dir: str) -> Optional[str]:
    if not work_dir:
        return None
    thermal_dir = os.path.join(work_dir, THERMAL_FOLDER_NAME)
    if os.path.isdir(thermal_dir):
        return thermal_dir
    if os.path.basename(work_dir).upper() == THERMAL_FOLDER_NAME:
        return work_dir
    return None


def build_cit_index(work_dir: str) -> dict:
    """
    Indexe toutes les mesures CIT d'un dossier LT en une passe
    (évite de rouvrir les mêmes .xlsm pour chaque SN).
    """
    thermal_dir = _resolve_thermal_dir(work_dir)
    if not thermal_dir:
        return {}

    index = {}
    xlsm_files = _list_thermal_files(thermal_dir, (".xlsm", ".xlsx"))
    Log.Message(f"[CIT] Indexation de {len(xlsm_files)} fichier(s) .xlsm/.xlsx...")
    for path in xlsm_files:
        file_index = _parse_xlsm_file_to_index(path)
        for key, value in file_index.items():
            if key not in index:
                index[key] = value

    acceptance_files = [
        p
        for p in _list_thermal_files(thermal_dir, (".xls",))
        if "acceptance" in os.path.basename(p).lower()
    ]
    if acceptance_files:
        Log.Message(f"[CIT] Indexation de {len(acceptance_files)} fichier(s) acceptance...")
        for path in acceptance_files:
            file_index = _parse_acceptance_file_to_index(path)
            for key, value in file_index.items():
                if key not in index:
                    index[key] = value

    Log.Message(f"[CIT] Index construit: {len(index)} mesure(s) de fin")
    return index


# Cache par dossier LT (évite double indexation dans la même session)
_cit_index_cache: dict = {}


def get_cit_index(work_dir: str) -> dict:
    cache_key = os.path.normcase(os.path.abspath(work_dir or ""))
    if cache_key not in _cit_index_cache:
        _cit_index_cache[cache_key] = build_cit_index(work_dir)
    return _cit_index_cache[cache_key]


def clear_cit_index_cache():
    _cit_index_cache.clear()


def _read_from_xlsm_com(path: str, sn: str, ref: str) -> Optional[str]:
    """Fallback lecture data_traitee via Excel COM (si openpyxl absent)."""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return None

    app = None
    wb = None
    try:
        try:
            pythoncom.CoInitialize()
        except Exception:
            pass
        try:
            app = win32com.client.DispatchEx("Excel.Application")
        except Exception:
            app = win32com.client.Dispatch("Excel.Application")
        try:
            app.Visible = 0
            app.DisplayAlerts = 0
        except Exception:
            pass

        wb = app.Workbooks.Open(path, ReadOnly=True)
        try:
            ws = wb.Sheets(DATA_SHEET_NAME)
        except Exception:
            return None

        used_cols = ws.UsedRange.Columns.Count
        used_rows = ws.UsedRange.Rows.Count
        header_row = HEADER_ROW_INDEX + 1  # Excel 1-based

        target_col = None
        header_label = None
        for col in range(1, used_cols + 1):
            header = ws.Cells(header_row, col).Value
            if _header_matches(header, sn, ref):
                target_col = col
                header_label = header
                break

        if target_col is None:
            return None

        raw = None
        for row in range(used_rows, header_row, -1):
            cell = ws.Cells(row, target_col).Value
            if _format_cit_value(cell) is not None:
                raw = cell
                break

        formatted = _format_cit_value(raw)
        if formatted is not None:
            Log.Message(
                f"[CIT] (Excel COM) {os.path.basename(path)} / {header_label} -> {formatted} "
                f"(SN{_normalize_sn(sn)})"
            )
        return formatted
    except Exception as e:
        Log.Warning(f"[CIT] Erreur lecture COM {os.path.basename(path)}: {e}")
        return None
    finally:
        try:
            if wb is not None:
                wb.Close(False)
        except Exception:
            pass
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass


def _read_from_xlsm(path: str, sn: str, ref: str) -> Optional[str]:
    try:
        import openpyxl
    except ImportError:
        Log.Warning("[CIT] openpyxl indisponible, lecture .xlsm via Excel COM")
        return _read_from_xlsm_com(path, sn, ref)

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        Log.Warning(f"[CIT] openpyxl échec sur {os.path.basename(path)}: {e}, repli Excel COM")
        return _read_from_xlsm_com(path, sn, ref)

    try:
        if DATA_SHEET_NAME not in wb.sheetnames:
            return None
        ws = wb[DATA_SHEET_NAME]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) <= HEADER_ROW_INDEX:
            return None

        headers = rows[HEADER_ROW_INDEX]
        data_rows = rows[HEADER_ROW_INDEX + 1 :]

        for col_index, header in enumerate(headers):
            if not _header_matches(header, sn, ref):
                continue
            raw = _find_last_numeric_in_column(data_rows, col_index)
            formatted = _format_cit_value(raw)
            if formatted is not None:
                try:
                    Log.Message(
                        f"[CIT] {os.path.basename(path)} / {header} -> {formatted} "
                        f"(SN{_normalize_sn(sn)})"
                    )
                except Exception:
                    pass
                return formatted
        return None
    except Exception as e:
        Log.Warning(f"[CIT] Erreur lecture {os.path.basename(path)}: {e}")
        return None
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _is_dashed_sn_token(token: str) -> bool:
    """True pour un SN OHDOP du type 25-20-01 (3 groupes de 2 chiffres)."""
    parts = token.split("-")
    return len(parts) == 3 and all(len(p) == 2 and p.isdigit() for p in parts)


def _parse_acceptance_filename_channels(filename: str) -> list:
    """
    Parse un nom OptoTest du type :
      'oha acceptance  23.157 142-143-144-145-vide-146-....XLS'
      'ohdop acceptance 25-20-01 25-20-02 25-20-10.XLS'
    Retourne une liste indexée 0 = CH1 : {'sn': '...', 'ref_hint': '...'} ou None si vide.
    """
    base = os.path.splitext(os.path.basename(filename))[0]
    rest = re.sub(r"(?i)^(?:oha|ohdop)\s*acceptance\s*", "", base).strip()
    if not rest:
        return []

    channels = []
    current_ref = ""
    tokens = re.split(r"\s+", rest)

    for token in tokens:
        token = token.strip().strip("-")
        if not token:
            continue
        if re.match(r"(?i)^(?:[A-Z°-]*\d+\.\d+[A-Z]?)$", token):
            current_ref = token
            continue
        if _is_dashed_sn_token(token):
            channels.append({"sn": _normalize_sn(token), "ref_hint": current_ref})
            continue
        if "-" in token:
            for part in token.split("-"):
                part = part.strip()
                if not part or part.lower() in ("vide", "empty", "na", "n/a"):
                    channels.append(None)
                elif re.match(r"(?i)^(?:[A-Z°-]*\d+\.\d+[A-Z]?)$", part):
                    current_ref = part
                elif re.match(r"^\d+(?:-\d+)*$", part):
                    channels.append({"sn": _normalize_sn(part), "ref_hint": current_ref})
            continue
        if token.lower() in ("vide", "empty", "na", "n/a"):
            channels.append(None)
        elif re.match(r"^\d+(?:-\d+)*$", token):
            channels.append({"sn": _normalize_sn(token), "ref_hint": current_ref})

    return channels


def _read_from_acceptance_xls(path: str, sn: str, ref: str) -> Optional[str]:
    """Fallback OptoTest .XLS via win32com."""
    channels = _parse_acceptance_filename_channels(path)
    if not channels:
        return None

    sn_norm = _normalize_sn(sn)
    ref_norm = _normalize_ref(ref)
    target_col = None  # 1-based Excel column for CH

    for idx, ch in enumerate(channels):
        if ch is None:
            continue
        if ch["sn"] != sn_norm:
            continue
        hint = _normalize_ref(ch.get("ref_hint") or "")
        if ref_norm and hint and not (ref_norm in hint or hint in ref_norm):
            continue
        # Header row has CH1 at column 6 (Excel col F) based on sample layout:
        # col2=Date, col3=Time, col4=Temp, col5=Wavelength, col6=CH1
        target_col = 6 + idx
        break

    if target_col is None:
        return None

    try:
        import pythoncom
        import win32com.client
    except ImportError:
        Log.Warning("[CIT] win32com indisponible pour lire les fichiers acceptance .XLS")
        return None

    app = None
    wb = None
    try:
        try:
            pythoncom.CoInitialize()
        except Exception:
            pass
        try:
            app = win32com.client.DispatchEx("Excel.Application")
        except Exception:
            app = win32com.client.Dispatch("Excel.Application")
        try:
            app.Visible = 0
            app.DisplayAlerts = 0
        except Exception:
            pass

        wb = app.Workbooks.Open(path, ReadOnly=True)
        ws = wb.Sheets(1)
        used_rows = ws.UsedRange.Rows.Count

        raw = None
        for row in range(used_rows, 6, -1):  # data starts ~row 7
            cell = ws.Cells(row, target_col).Value
            if _format_cit_value(cell) is not None:
                raw = cell
                break

        formatted = _format_cit_value(raw)
        if formatted is not None:
            Log.Message(
                f"[CIT] acceptance {os.path.basename(path)} CH{target_col - 5} "
                f"-> {formatted} (SN{sn_norm})"
            )
        return formatted
    except Exception as e:
        Log.Warning(f"[CIT] Erreur lecture acceptance {os.path.basename(path)}: {e}")
        return None
    finally:
        try:
            if wb is not None:
                wb.Close(False)
        except Exception:
            pass
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass


def fetch_fin_cit_for_sn(
    work_dir: str, sn: str, ref_sedi: str = "", index: dict = None
) -> Optional[str]:
    """
    Cherche la mesure de fin CIT pour un SN dans work_dir/THERMAL_CYCLING.
    Retourne une chaîne formatée (ex. '0.052') ou None.
    """
    if not work_dir or not sn:
        return None

    if index is None:
        index = get_cit_index(work_dir)

    value = _lookup_cit_in_index(index, sn, ref_sedi)
    if value is not None:
        Log.Verbose(
            f"[CIT] SN{_normalize_sn(sn)} (ref={ref_sedi or '-'}) -> {value}"
        )
        return value

    if not _resolve_thermal_dir(work_dir):
        Log.Verbose(f"[CIT] Dossier THERMAL_CYCLING introuvable sous {work_dir}")

    Log.Verbose(
        f"[CIT] Aucune mesure de fin pour SN{_normalize_sn(sn)} "
        f"(ref={ref_sedi or '-'})"
    )
    return None


def container_has_cit_value(container) -> bool:
    tags = getattr(container, "tagAndValues", {}) or {}
    for key in CIT_TAG_ALIASES:
        value = tags.get(key)
        if value is not None and str(value).strip() != "":
            return True
    return False


def apply_fin_cit_to_container(container, work_dir: str) -> bool:
    """
    Remplit CIT/FIN_CIT/fin_cit si vides. Retourne True si une valeur a été écrite.
    """
    if container is None or container_has_cit_value(container):
        return False

    tags = getattr(container, "tagAndValues", None)
    if tags is None:
        container.tagAndValues = {}
        tags = container.tagAndValues

    ref = (
        tags.get("REF_SEDI")
        or tags.get("**REF_SEDI**")
        or tags.get("REF_CLIENT")
        or tags.get("**REF_CLIENT**")
        or tags.get("NUMPLAN")
        or tags.get("**NUMPLAN**")
        or ""
    )
    sn = getattr(container, "SN", "") or tags.get("SN") or ""

    value = fetch_fin_cit_for_sn(work_dir, sn, str(ref))
    if value is None:
        return False

    for key in CIT_TAG_ALIASES:
        tags[key] = value

    undefined = getattr(container, "undefinedValues", None)
    if isinstance(undefined, list):
        for key in CIT_TAG_ALIASES:
            while key in undefined:
                undefined.remove(key)

    Log.Message(f"[CIT] SN{sn} : mesure de fin auto-remplie = {value}")
    return True


def apply_fin_cit_to_containers(containers, work_dir: str) -> int:
    """Applique le remplissage auto sur une liste de containers. Retourne le nombre remplis."""
    if not containers:
        return 0

    # Une seule indexation pour tous les SN (200+ mesures)
    cit_index = get_cit_index(work_dir)
    filled = 0
    missing = 0
    for container in containers:
        try:
            if container_has_cit_value(container):
                continue
            tags = getattr(container, "tagAndValues", None) or {}
            ref = (
                tags.get("REF_SEDI")
                or tags.get("**REF_SEDI**")
                or tags.get("REF_CLIENT")
                or tags.get("**REF_CLIENT**")
                or tags.get("NUMPLAN")
                or tags.get("**NUMPLAN**")
                or ""
            )
            sn = getattr(container, "SN", "") or tags.get("SN") or ""
            value = fetch_fin_cit_for_sn(work_dir, sn, str(ref), index=cit_index)
            if value is None:
                missing += 1
                continue
            if tags is not container.tagAndValues:
                container.tagAndValues = tags
            for key in CIT_TAG_ALIASES:
                tags[key] = value
            undefined = getattr(container, "undefinedValues", None)
            if isinstance(undefined, list):
                for key in CIT_TAG_ALIASES:
                    while key in undefined:
                        undefined.remove(key)
            filled += 1
            Log.Verbose(f"[CIT] SN{sn} : mesure de fin auto-remplie = {value}")
        except Exception as e:
            Log.Warning(f"[CIT] Échec auto-remplissage pour SN{getattr(container, 'SN', '?')}: {e}")

    if filled or missing:
        Log.Message(
            f"[CIT] {filled} mesure(s) auto-remplie(s)"
            + (f", {missing} SN sans donnée CIT" if missing else "")
        )
    return filled
